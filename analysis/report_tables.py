"""
Recompute every number the report quotes, from the result files and the
success labels.

The result files and labels live on the `results-archive` branch. Bring them
into the (git-ignored) results/ folder first, then run from the project root:

    git restore --source=origin/results-archive --worktree -- results/
    pip install openpyxl
    python3 analysis/report_tables.py

A success counts as "genuine" when the case-by-case review judged that the
paraphrase asks the same question and the victim's answer is wrong
(results/labels/: one sheet per main-study run, plus
earlier_runs_genuine.json for the smaller runs).
"""

import argparse
import glob
import json
import os
import re
from collections import Counter

from openpyxl import load_workbook

MAIN = "*q96_p5_r8_evolutionary__main_study_combined.json"
MAIN_B1 = "*semantic_preserve*q96_p40_r1*20261001_000152.json"
BASELINES_16 = [  # name, file pattern
    ("B1 plain paraphrasing, one shot", "*semantic_preserve*q16_p40_r1*20260928_184240.json"),
    ("B2 attacker prompt, one shot", "*minimal_edit*q16_p40_r1*20260928_185644.json"),
    ("B3 reflective, 8 rounds", "*minimal_edit*q16_p5_r8_reflective__20260928_191017.json"),
    ("B4 evolutionary, 8 rounds", "*minimal_edit*q16_p5_r8_evolutionary__20260928_193042.json"),
]
REASONING_22 = "*victim_qwen3_4b__*q22_p5_r8*20260927_053549.json"
INSTRUCT_16 = "*victim_qwen3_4b_instruct__*q16_p5_r8*20260927_142038.json"
SHEETS = {
    MAIN: "main_study_full_method_review.xlsx",
    MAIN_B1: "main_study_B1_review.xlsx",
}
CONFIDENCE_GROUPS = ("certain", "mostly", "unsure")


def load(results_dir, pattern):
    paths = glob.glob(os.path.join(results_dir, pattern))
    if len(paths) != 1:
        raise SystemExit(f"expected one file for {pattern} in {results_dir}, found {len(paths)}")
    with open(paths[0], encoding="utf-8") as fh:
        return json.load(fh)


def successes(run):
    """(question_id, success record) for every success in a run."""
    return [(r["question_id"], s) for r in run["results"] for s in r["successful_attacks"]]


def sheet_labels(results_dir, sheet):
    """{(question_id, paraphrase): label} from a review sheet."""
    ws = load_workbook(os.path.join(results_dir, "labels", sheet), read_only=True)["Successes"]
    return {(row[1], row[6]): row[11] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]}


def earlier_genuine(results_dir):
    with open(os.path.join(results_dir, "labels", "earlier_runs_genuine.json"), encoding="utf-8") as fh:
        return {k: v for k, v in json.load(fh).items() if not k.startswith("_")}


def genuine_from_list(run, spec):
    return [(q, s) for q, s in successes(run)
            if q in spec["all_successes_of"] or s["paraphrase"] in spec["paraphrases"]]


def candidate_counts(run):
    """Status / rejection reason of every non-duplicate candidate."""
    c = Counter()
    for r in run["results"]:
        for rd in r["rounds"]:
            for cand in rd["candidates"]:
                if cand["status"] == "duplicate":
                    continue
                c[cand["status"] if cand["status"] != "rejected"
                  else "rejected: " + cand.get("rejected_by", "?")] += 1
    return c


def table(title, header, rows):
    print(f"\n{title}\n" + "-" * len(title))
    widths = [max(len(str(x)) for x in col) for col in zip(header, *rows)]
    for line in [header] + rows:
        print("  ".join(str(x).ljust(w) for x, w in zip(line, widths)))


def pct(a, b):
    return f"{a / b:.1%}" if b else "-"


def main_study(results_dir):
    runs = {"Full method": (MAIN, load(results_dir, MAIN)),
            "B1 plain paraphrasing": (MAIN_B1, load(results_dir, MAIN_B1))}
    rows, label_rows, group_rows, per_question = [], [], [], {}
    for name, (pattern, run) in runs.items():
        labels = sheet_labels(results_dir, SHEETS[pattern])
        succ = successes(run)
        lab = [labels[(q, s["paraphrase"])] for q, s in succ]
        gen = [(q, s) for (q, s), l in zip(succ, lab) if l == "genuine"]
        gen_conf = [(q, s) for q, s in gen if s.get("confirmed_success")]
        tested = sum(r["n_valid_victim_queries"] for r in run["results"])
        attacked = sum(1 for r in run["results"] if r["rounds"])
        q_gc = {q for q, _ in gen_conf}
        per_question[name] = Counter(q for q, _ in gen_conf)
        rows.append([name, attacked, tested, len(succ),
                     sum(1 for _, s in succ if s.get("robust_success")),
                     sum(1 for _, s in succ if s.get("confirmed_success")),
                     len(gen), len(gen_conf), f"{len(q_gc)} ({pct(len(q_gc), attacked)})",
                     pct(len(gen_conf), tested)])
        lc = Counter(lab)
        label_rows.append([name] + [lc[k] for k in
                                    ("genuine", "drift", "borderline", "judge error", "garbled", "typo")])
        group = {r["question_id"]: r.get("confidence_group") for r in run["results"] if r["rounds"]}
        for g in CONFIDENCE_GROUPS:
            qs = [q for q, gg in group.items() if gg == g]
            gc = [q for q, _ in gen_conf if group.get(q) == g]
            group_rows.append([name, g, len(qs), len(gc), f"{len(set(gc))} ({pct(len(set(gc)), len(qs))})"])

    table("Main study: 96 questions, victim qwen3:4b-instruct",
          ["Run", "Attacked q", "Paraphrases tested", "Successes", "Robust", "Confirmed",
           "Genuine", "Genuine+confirmed", "Questions with genuine+confirmed", "Genuine+confirmed per tested"],
          rows)
    table("Labels of all successes (main study)",
          ["Run", "genuine", "drift", "borderline", "judge error", "garbled", "typo"], label_rows)
    table("By the victim's confidence on the original (20 samples)",
          ["Run", "Group", "Questions", "Genuine+confirmed", "Questions with genuine+confirmed"], group_rows)

    full, b1 = per_question["Full method"], per_question["B1 plain paraphrasing"]
    both = set(full) & set(b1)
    print(f"\nQuestions with a genuine+confirmed success: full method {len(full)}, B1 {len(b1)}, "
          f"both {len(both)}, only full method {len(set(full) - both)}, only B1 {len(set(b1) - both)}")
    print(f"Genuine+confirmed per such question: full method {sum(full.values()) / len(full):.1f}, "
          f"B1 {sum(b1.values()) / len(b1):.1f}")

    table("What happened to every generated paraphrase (main study, duplicates excluded)",
          ["Outcome"] + list(runs),
          [[k] + [candidate_counts(run)[k] for _, run in runs.values()]
           for k in sorted(set().union(*(candidate_counts(run) for _, run in runs.values())))])


def baselines_16(results_dir):
    lists = earlier_genuine(results_dir)
    rows = []
    for name, pattern in BASELINES_16:
        run = load(results_dir, pattern)
        ts = re.search(r"\d{8}_\d{6}", pattern).group()
        gen = genuine_from_list(run, lists[ts])
        cc = candidate_counts(run)
        generated = sum(cc.values())
        rows.append([name, sum(1 for r in run["results"] if r["rounds"]),
                     sum(r["n_valid_victim_queries"] for r in run["results"]), len(successes(run)),
                     len(gen), len({q for q, _ in gen}),
                     pct(cc["rejected: too_different"], generated)])
    table("Baselines on the 16-question set (A100, budget 40 paraphrases per question)",
          ["Run", "Attacked q", "Paraphrases tested", "Successes", "Genuine",
           "Questions with genuine", "Rejected as too different"], rows)


def reasoning_vs_instruct(results_dir):
    lists = earlier_genuine(results_dir)
    instruct = load(results_dir, INSTRUCT_16)
    reasoning = load(results_dir, REASONING_22)
    shared = {r["question_id"] for r in instruct["results"]}
    rows = []
    for name, run, ts in (("qwen3:4b (reasoning)", reasoning, "20260927_053549"),
                          ("qwen3:4b-instruct", instruct, "20260927_142038")):
        succ = [(q, s) for q, s in successes(run) if q in shared]
        gen = [(q, s) for q, s in genuine_from_list(run, lists[ts]) if q in shared]
        rows.append([name, len(shared), len(succ), len(gen), len({q for q, _ in gen})])
    table("Reasoning vs non-reasoning victim, on the 16 questions both answer correctly",
          ["Victim", "Questions", "Successes", "Genuine", "Questions with genuine"], rows)
    gen22 = genuine_from_list(reasoning, lists["20260927_053549"])
    print(f"Reasoning victim on all 22 questions: {len(successes(reasoning))} successes, "
          f"{len(gen22)} genuine on {len({q for q, _ in gen22})} questions")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--results-dir", default="results")
    args = parser.parse_args()
    main_study(args.results_dir)
    baselines_16(args.results_dir)
    reasoning_vs_instruct(args.results_dir)


if __name__ == "__main__":
    main()
